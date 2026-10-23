import io
import json
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock

import docx
import openpyxl  # type: ignore[import-untyped]
import pptx
import pytest
from docling_core.types.doc.document import DoclingDocument

from edda.drivers import DriverError, import_document
from edda.drivers.document import Converted, _convert, _office, converter_name, models_dir, paged_markdown
from edda.kernel import Kernel
from edda.models import AddedSource, SourceSummary

RECORDED = Path(__file__).parents[1] / "fixtures" / "docling"
PDF = b"%PDF-1.7\n..."


def recorded(name: str) -> tuple[DoclingDocument, list[int]]:
    """Docling's real output for a fixture PDF, recorded with the approved models (tests never run them)."""
    record = json.loads((RECORDED / f"{name}.json").read_text(encoding="utf-8"))
    return DoclingDocument.model_validate(record["document"]), record["pages_without_text"]


@pytest.fixture
def kernel() -> MagicMock:
    fake = MagicMock(spec=Kernel)
    count = iter(range(1, 100))

    def add_source(data: bytes, **meta: Any) -> AddedSource:
        summary = SourceSummary(hash=f"{next(count):02d}" * 32, title=meta["title"], driver=meta["driver"])
        return AddedSource(source=summary.model_copy(update={"derived_from": meta.get("derived_from")}), added=True)

    fake.add_source.side_effect = add_source
    return fake


def test_pages_of_a_recorded_conversion_are_marked() -> None:
    document, without_text = recorded("text")
    text = paged_markdown(document)
    assert text.startswith("<!-- page 1 -->\n\nLantern protocol Lantern is a protocol")
    assert "<!-- page 2 -->\n\nHistory\n\nThe first version was drafted in 2019" in text
    assert without_text == []
    scan, without_text = recorded("scan")
    assert paged_markdown(scan) == "<!-- page 1 -->\n\nScannedmemo\n\nThe field station syncs once a day at noon."
    assert without_text == [1]


def test_an_empty_page_keeps_its_marker() -> None:
    document = MagicMock(pages={2: None, 1: None})
    document.export_to_markdown.side_effect = lambda page_no: "" if page_no == 2 else "One."
    assert paged_markdown(document) == "<!-- page 1 -->\n\nOne.\n\n<!-- page 2 -->"


def test_stores_the_original_and_the_markdown_made_from_it(kernel: MagicMock) -> None:
    document, _ = recorded("text")
    convert = lambda data, ext: Converted(text=paged_markdown(document), converter="docling 2.96.1", method="layout")  # noqa: E731
    imported = import_document(kernel, PDF, filename="Lantern.PDF", convert=convert)
    original, derived = kernel.add_source.call_args_list
    assert original.args == (PDF,)
    assert original.kwargs == {"extension": "pdf", "title": "Lantern", "driver": "document", "by": "owner"}
    assert derived.args[0].decode().startswith("<!-- page 1 -->")
    assert derived.kwargs["derived_from"] == imported.original.source.hash
    assert (derived.kwargs["converter"], derived.kwargs["method"], derived.kwargs["extension"]) == (
        "docling 2.96.1",
        "layout",
        "md",
    )
    assert imported.derived.source.derived_from == imported.original.source.hash


def test_titles_come_from_the_first_heading_or_the_caller(kernel: MagicMock) -> None:
    convert = lambda data, ext: Converted(text="# Field notes\n\nText.", converter="c", method="native")  # noqa: E731
    import_document(kernel, b"PK\x03\x04rest", filename="n.docx", convert=convert)
    assert kernel.add_source.call_args.kwargs["title"] == "Field notes"
    import_document(kernel, b"PK\x03\x04rest", filename="n.docx", convert=convert, title="Mine")
    assert kernel.add_source.call_args.kwargs["title"] == "Mine"


@pytest.mark.parametrize(
    ("data", "filename", "message"),
    [
        (b"", "a.pdf", "the file is empty"),
        (b"hello", "a.txt", "not a document this importer reads"),
        (b"hello", "a.pdf", r"does not look like a \.pdf file"),
        (b"%PDF-1.7", "a.docx", r"does not look like a \.docx file"),
        (b"\x89PNG", "a.jpg", r"does not look like a \.jpg file"),
    ],
)
def test_refuses_files_it_cannot_read_and_stores_nothing(
    kernel: MagicMock, data: bytes, filename: str, message: str
) -> None:
    with pytest.raises(DriverError, match=message):
        import_document(kernel, data, filename=filename, convert=lambda d, e: pytest.fail("converted"))
    kernel.add_source.assert_not_called()


def test_a_document_with_no_text_is_refused(kernel: MagicMock) -> None:
    with pytest.raises(DriverError, match=r"found no text in blank\.png"):
        import_document(
            kernel,
            b"\x89PNG...",
            filename="blank.png",
            convert=lambda d, e: Converted(text=" \n", converter="c", method="ocr"),
        )
    kernel.add_source.assert_not_called()


def _docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Lantern protocol", 1)
    doc.add_paragraph("Lantern batches requests over slow links.")
    table = doc.add_table(rows=2, cols=2)
    for (r, c), text in {(0, 0): "k1", (0, 1): "b", (1, 0): "1.2", (1, 1): "0.75"}.items():
        table.cell(r, c).text = text
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _pptx() -> bytes:
    deck = pptx.Presentation()
    for title, body in [("Field trials", "Batching saved time."), ("Results", "Two sites synced daily.")]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def _xlsx() -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Trials"
    for row in (["site", "syncs per day"], ["North", 1]):
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_office_files_convert_natively_with_markitdown() -> None:
    word = _office(_docx(), ".docx")
    assert (word.method, word.converter.split()[0]) == ("native", "markitdown")
    assert word.text.startswith("# Lantern protocol\n\nLantern batches requests over slow links.")
    assert "| 1.2 | 0.75 |" in word.text
    slides = _office(_pptx(), ".pptx").text
    assert "<!-- page 1 -->\n# Field trials" in slides and "<!-- page 2 -->\n# Results" in slides
    assert "| North | 1 |" in _convert(_xlsx(), ".xlsx").text


def test_a_damaged_office_file_is_a_driver_error_not_its_bytes_as_text() -> None:
    with pytest.raises(DriverError, match=r"the \.docx file is damaged: it has no word/document\.xml"):
        _office(b"PK\x03\x04 not really a zip", ".docx")
    with pytest.raises(DriverError, match=r"the \.pptx file is damaged"):
        _office(_docx(), ".pptx")


def test_a_markitdown_failure_is_a_driver_error(monkeypatch: pytest.MonkeyPatch) -> None:
    import markitdown

    def fail(*_: object, **__: object) -> None:
        raise markitdown.FileConversionException("bad table")

    monkeypatch.setattr(markitdown.MarkItDown, "convert_stream", fail)
    with pytest.raises(DriverError, match=r"could not read the \.docx file: bad table"):
        _office(_docx(), ".docx")


def test_password_protected_and_old_office_files_get_a_clear_refusal(kernel: MagicMock) -> None:
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest"
    with pytest.raises(DriverError, match=r"password-protected, or an old binary Office file; save it as \.xlsx"):
        import_document(kernel, ole, filename="budget.xlsx")


def test_the_pdf_lane_needs_the_models(home: Path) -> None:
    assert models_dir() == home / ".cache" / "edda" / "docling"
    with pytest.raises(DriverError, match="run `pixi run doc-models`"):
        _convert(PDF, ".pdf")


def test_names_each_converter_with_its_installed_version() -> None:
    from importlib.metadata import version

    assert converter_name("docling") == f"docling {version('docling-slim')}"
    assert converter_name("markitdown") == f"markitdown {version('markitdown')}"
